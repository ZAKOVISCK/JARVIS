"""Attach NUMREG only after every existing part row matches the source export."""
import hashlib, json, re, sys, tempfile
from pathlib import Path
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[1]; target=root/'public/data'
source=Path(sys.argv[1]); docs={}; count=0
clean=lambda v:' '.join(str(v or '').split())
w=load_workbook(source,read_only=True,data_only=True); rows=w.active.iter_rows(values_only=True)
h=[clean(v) for v in next(rows)]; ix={v:i for i,v in enumerate(h)}
for n,row in enumerate(rows,2):
    code=clean(row[ix['CARRO']]); code=str(int(code)) if code.isdigit() else ''
    if len(code)==5 and code.startswith('55'): code=str(int(code)-55000)
    if not code.isdigit() or not 1<=int(code)<=559: continue
    rawdate=row[ix['DATA']]; date=rawdate.strftime('%Y-%m-%d') if hasattr(rawdate,'strftime') else '-'.join(clean(rawdate).split('/')[::-1])
    qty=row[ix['QUANT']]; qty=float(qty) if qty is not None and str(qty).strip() else None
    doc=clean(row[ix['NUMREG']]); doc=doc if doc.isdigit() else None
    docs[n]=(code,date,clean(row[ix['MATERIAL']]),qty,doc)
w.close()
with tempfile.TemporaryDirectory(prefix='jarvis-evidence-') as tmp:
    stage=Path(tmp); fingerprints={}; matched=0
    for path in sorted(target.glob('*.json')):
        if not path.stem.isdigit(): continue
        data=json.loads(path.read_text())
        for part in data['parts']:
            original=docs.get(part[6]); expected=(str(int(data['code'])),part[0],part[3],part[4])
            if not original or original[:4]!=expected:
                raise RuntimeError(f'Origem divergente: prefixo {data["code"]}, linha {part[6]}')
            part[8:]=[original[4],None]; matched+=1
            if original[4]: count+=1
        output=json.dumps(data,ensure_ascii=False,separators=(',',':')).encode()
        (stage/path.name).write_bytes(output); fingerprints[path.stem]=hashlib.sha256(output).hexdigest()
    if matched!=len(docs): raise RuntimeError(f'Contagem divergente {matched}/{len(docs)}')
    identity=hashlib.sha256(json.dumps(fingerprints,sort_keys=True).encode()).hexdigest()
    manifest={'version':identity,'file':source.name,'sourceSha256':hashlib.sha256(source.read_bytes()).hexdigest(),'contract':'stock-extract-with-document','documentColumn':'NUMREG','timeAvailable':False,'records':matched,'withDocument':count,'vehicleHashes':fingerprints}
    for p in stage.glob('*.json'): (target/p.name).write_bytes(p.read_bytes())
    (target/'source-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,separators=(',',':')))
print(json.dumps({'matched':matched,'withDocument':count,'version':identity}))
