"""Build private Site lookup shards from the maintenance and fleet exports.

Usage: python scripts/prepare-fleet-data.py LISTAGEM.xlsx EXTRATO.xlsx FROTA.xlsx
Only fields needed for maintenance lookup are retained. Driver names and costs
are deliberately excluded. Source line numbers preserve traceability.
"""
from __future__ import annotations

import json
import hashlib
import re
import sys
import tempfile
from collections import Counter, OrderedDict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "public" / "data"


def numeric_code(value):
    raw = str(value or "").strip()
    if not re.fullmatch(r"\d+", raw):
        return None
    return str(int(raw))


def vehicle_code(value):
    """Normalize either 001..559 or the displayed 55001..55559 prefix."""
    raw = re.sub(r"\s+", "", str(value or ""))
    if not re.fullmatch(r"\d+", raw):
        return None
    number = int(raw)
    if len(raw) == 5 and raw.startswith("55"):
        number -= 55000
    if 1 <= number <= 559:
        return str(number)
    return None


def order(value):
    raw = str(value or "").strip()
    if not raw or not re.fullmatch(r"\d+", raw):
        return ""
    return str(int(raw))


def clean(value):
    return " ".join(str(value or "").split())


def family(value):
    """Drop the stock class number and keep its human-readable denomination."""
    label = re.sub(r"^\d+\s*", "", clean(value)).strip()
    return label.title()


def iso(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    raw = clean(value)
    if re.fullmatch(r"\d{2}/\d{2}/\d{4}", raw):
        return "-".join(reversed(raw.split("/")))
    return raw[:10]


class Stage:
    def __init__(self, directory):
        self.directory = Path(directory)
        self.handles = OrderedDict()
        self.codes = set()

    def append(self, kind, vehicle, record):
        self.codes.add(vehicle)
        key = (kind, vehicle)
        if key not in self.handles:
            if len(self.handles) >= 64:
                _, old = self.handles.popitem(last=False)
                old.close()
            self.handles[key] = (self.directory / f"{kind}-{vehicle}.jsonl").open("a", encoding="utf-8")
        handle = self.handles[key]
        self.handles.move_to_end(key)
        handle.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n")

    def close(self):
        for handle in self.handles.values():
            handle.close()
        self.handles.clear()


def read_lines(path):
    if not path.exists():
        return []
    with path.open(encoding="utf-8") as stream:
        return [json.loads(line) for line in stream]


def load_fleet(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "GERAL" not in workbook.sheetnames:
        raise ValueError("A planilha de frota precisa conter a aba GERAL")
    sheet = workbook["GERAL"]
    header = [clean(x) for x in next(sheet.values)]
    ix = {name: i for i, name in enumerate(header)}
    required = {"PREFIXO", "ANO MODELO", "MARCA CHASSI", "MODELO CHASSI", "PLACA"}
    missing = required - set(ix)
    if missing:
        raise ValueError(f"Colunas ausentes na frota: {', '.join(sorted(missing))}")
    vehicles = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        car = vehicle_code(row[ix["PREFIXO"]])
        if car is None:
            continue
        if car in vehicles:
            raise ValueError(f"Prefixo duplicado na frota: {car}")
        year = row[ix["ANO MODELO"]]
        year = int(year) if isinstance(year, (int, float)) else clean(year) or None
        vehicles[car] = {
            "year": year,
            "brand": clean(row[ix["MARCA CHASSI"]]) or None,
            "model": clean(row[ix["MODELO CHASSI"]]) or None,
            "plate": clean(row[ix["PLACA"]]).upper() or None,
        }
    workbook.close()
    return vehicles


def build(listing, extract, fleet_source):
    OUTPUT.mkdir(parents=True, exist_ok=True)
    counts = Counter()
    fleet_metadata = load_fleet(fleet_source)
    with tempfile.TemporaryDirectory(prefix="fleet-stage-") as directory:
        stage = Stage(directory)
        ready = Path(directory) / "ready"
        ready.mkdir()
        fingerprints = {}
        for kind, path in (("s", listing), ("p", extract)):
            sheet = load_workbook(path, read_only=True, data_only=True).active
            header = [clean(x) for x in next(sheet.values)]
            ix = {name: i for i, name in enumerate(header)}
            for rownum, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                car = vehicle_code(row[ix["CARRO"]])
                if car is None:
                    counts[f"skipped_{kind}"] += 1
                    continue
                d, os = iso(row[ix["DATA"]]), order(row[ix["O.S."]])
                if kind == "s":
                    # date, FICHA, MEC code, action, occurrence, alleged defect, source row, OS (join key)
                    record = [d, order(row[ix["FICHA"]]), numeric_code(row[ix["MEC"]]), clean(row[ix["PROVIDENCIAS"]]),
                              clean(row[ix["OCORRENCIA"]]), clean(row[ix["DEFEITO ALEGADO"]]), rownum, os]
                else:
                    # date of issue, OS, chassis KM, material, quantity, FUNC code, source row, family
                    km = row[ix["KM CHASSI"]]
                    km = int(km) if isinstance(km, (float, int)) and km >= 0 else None
                    qty = row[ix["QUANT"]]
                    if isinstance(qty, (int, float)):
                        counts["positive_p" if qty > 0 else "negative_p" if qty < 0 else "zero_p"] += 1
                    record = [d, os, km, clean(row[ix["MATERIAL"]]), qty, numeric_code(row[ix["FUNC."]]), rownum,
                              family(row[ix["CLASSE/DENOMINACAO"]]), order(row[ix["NUMREG"]]) or None, None]
                stage.append(kind, car, record)
                counts[kind] += 1
                if not os:
                    counts[f"without_os_{kind}"] += 1
                if rownum % 100000 == 0:
                    print(f"{path.name}: {rownum - 1:,} rows", flush=True)
            stage.close()

        fleet = []
        for car in sorted(stage.codes | set(fleet_metadata), key=lambda c: int(c)):
            metadata = fleet_metadata.get(car, {"year": None, "brand": None, "model": None, "plate": None})
            services = read_lines(Path(directory) / f"s-{car}.jsonl")
            parts = read_lines(Path(directory) / f"p-{car}.jsonl")
            services.sort(key=lambda r: (r[0], r[6]), reverse=True)
            parts.sort(key=lambda r: (r[0], r[6]), reverse=True)
            service_os = {r[7] for r in services if r[7]}
            counts["parts_with_service_os"] += sum(1 for r in parts if r[1] in service_os)
            last = max((services[0][0] if services else ""), (parts[0][0] if parts else ""))
            fleet.append({"code": car, **metadata, "last": last,
                          "services": len(services), "parts": len(parts)})
            (ready / f"{car}.json").write_text(json.dumps({"code": car, **metadata,
                "services": services, "parts": parts}, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
            fingerprints[car] = hashlib.sha256((ready / f"{car}.json").read_bytes()).hexdigest()
        manifest = {"sources": [listing.name, extract.name, fleet_source.name], "counts": dict(counts), "fleet": fleet}
        (ready / "summary.json").write_text(json.dumps(manifest, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        identity = hashlib.sha256(json.dumps(fingerprints, sort_keys=True).encode()).hexdigest()
        provenance = {"version": identity, "file": extract.name, "sourceSha256": hashlib.sha256(extract.read_bytes()).hexdigest(),
                      "contract": "stock-extract-with-document", "documentColumn": "NUMREG", "timeAvailable": False,
                      "records": counts["p"], "vehicleHashes": fingerprints}
        (ready / "source-manifest.json").write_text(json.dumps(provenance, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        # All source parsing succeeds before replacing tracked output; never clear the dataset upfront.
        for prepared in ready.glob("*.json"):
            prepared.replace(OUTPUT / prepared.name)
        for previous in OUTPUT.glob("*.json"):
            if previous.stem.isdigit() and previous.stem not in fingerprints:
                previous.unlink()
        print("Rebuild derived catalog and administrative audit after source updates: npm run catalog:build && npm run audit", flush=True)
        print(json.dumps({"vehicles": len(fleet), "counts": dict(counts),
                          "bytes": sum(p.stat().st_size for p in OUTPUT.glob("*.json"))}, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit("Provide the LISTAGEM, EXTRATO and FROTA .xlsx paths")
    build(Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3]))
